"""Builds a реферат as a Word document formatted by ГОСТ 7.32-2017 from the
same ReferenceContent the PDF is built from: title page, СОДЕРЖАНИЕ as a real
TOC field (Word fills in page numbers on open), the text with built-in
Heading 1/Heading 2 styles, and СПИСОК ИСТОЧНИКОВ.
"""
import html
import re
from datetime import date, datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.reference.pdf_writer import ReferenceContent, TitlePageInfo

FONT_NAME = "Times New Roman"
BODY_STYLE = "Body Text"
TOC_INSTRUCTION = 'TOC \\o "1-2" \\h \\z \\u'
TOC_PLACEHOLDER = "Оглавление обновится при открытии документа"

_CENTER = WD_ALIGN_PARAGRAPH.CENTER
_LEFT = WD_ALIGN_PARAGRAPH.LEFT
_INDENT = Cm(1.25)
_SIGNATURE_COLUMN_WIDTHS = (Cm(6.5), Cm(3.5), Cm(6.5))
_CAPTION_SIZE = 10

_BR_TAG = re.compile(r"<br\s*/?>", re.IGNORECASE)
_ANY_TAG = re.compile(r"<[^>]*>")
_PARAGRAPH_BREAK = re.compile(r"\n\s*\n")
_WHITESPACE_RUN = re.compile(r"[ \t\r\n\x0b\x0c]+")
_XML_INVALID = re.compile(r"[\x00-\x08\x0e-\x1f\ufffe\uffff]")

# CT_Settings is a strict sequence; Word reports the file as damaged when
# updateFields is placed after any of these.
_SETTINGS_AFTER_UPDATE_FIELDS = (
    "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
    "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
    "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults",
    "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
)
_TC_PR_AFTER_BORDERS = (
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
)


def build_reference_docx(output_path: Path, content: ReferenceContent, font_size: int = 12) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)

    _setup_styles(doc, font_size)
    _add_title_page(doc, content.title_page)
    _add_table_of_contents(doc)
    _add_sections(doc, content.sections)
    _add_sources(doc, content.sources)
    _add_page_number_footer(section)
    _setup_settings(doc)
    _set_core_properties(doc, content.title_page)
    doc.save(str(output_path))


def _clean(text: str) -> str:
    return _WHITESPACE_RUN.sub(" ", _XML_INVALID.sub("", text)).strip()


def _paragraphs(text: str) -> list[str]:
    return [chunk for chunk in map(_clean, _PARAGRAPH_BREAK.split(text)) if chunk]


def _markup_lines(markup: str) -> list[str]:
    text = _ANY_TAG.sub("", _BR_TAG.sub("\n", markup)).replace("&nbsp;", " ")
    return [line for line in map(_clean, html.unescape(text).split("\n")) if line]


def _format(paragraph_format, **attributes) -> None:
    for name, value in attributes.items():
        setattr(paragraph_format, name, value)


def _set_font_family(rpr) -> None:
    fonts = rpr.get_or_add_rFonts()
    for theme_attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme_attribute), None)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), FONT_NAME)


def _setup_styles(doc, font_size: int) -> None:
    styles = doc.styles
    defaults = styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _set_font_family(defaults)
    defaults.find(qn("w:lang")).set(qn("w:val"), "ru-RU")

    normal = styles["Normal"]
    _set_font_family(normal.element.get_or_add_rPr())
    normal.font.size = Pt(font_size)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    _format(normal.paragraph_format, space_before=Pt(0), space_after=Pt(0), line_spacing=1.0)

    body = styles[BODY_STYLE]
    _format(
        body.paragraph_format,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=_INDENT,
        line_spacing=1.5, space_before=Pt(0), space_after=Pt(0),
    )

    for name in ("Heading 1", "Heading 2"):
        heading = styles[name]
        _set_font_family(heading.element.get_or_add_rPr())
        heading.font.size = Pt(font_size)
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.font.bold = True
        heading.font.italic = False
        heading.next_paragraph_style = body
        _format(
            heading.paragraph_format,
            alignment=_LEFT, first_line_indent=_INDENT, line_spacing=1.5,
            space_before=Pt(12), space_after=Pt(6), keep_with_next=True,
        )
    styles["Heading 2"].paragraph_format.left_indent = Cm(0)

    # builtin=True: a custom style named "toc 1" is ignored when Word builds the TOC.
    for level in (1, 2):
        toc = styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        toc.base_style = toc.next_paragraph_style = normal
        _format(toc.paragraph_format, line_spacing=1.5, left_indent=Cm(0.5 * (level - 1)))


def _fill(paragraph, text: str, align=_CENTER, bold: bool = False, size: int | None = None,
          space_before: int = 0):
    paragraph.alignment = align
    if space_before:
        paragraph.paragraph_format.space_before = Pt(space_before)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold or None
        if size:
            run.font.size = Pt(size)
    return paragraph


def _add_bottom_border(cell) -> None:
    bottom = OxmlElement("w:bottom")
    for attribute, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "0"), ("w:color", "000000")):
        bottom.set(qn(attribute), value)
    borders = OxmlElement("w:tcBorders")
    borders.append(bottom)
    cell._tc.get_or_add_tcPr().insert_element_before(borders, *_TC_PR_AFTER_BORDERS)


def _add_signature_table(doc, values: tuple[str, str, str], captions: tuple[str, str, str], first_align) -> None:
    table = doc.add_table(rows=2, cols=3)
    table.autofit = False
    for column, width in zip(table.columns, _SIGNATURE_COLUMN_WIDTHS):
        column.width = width
        for cell in column.cells:
            cell.width = width
    for cell, value, align in zip(table.rows[0].cells, values, (first_align, _CENTER, _CENTER)):
        _fill(cell.paragraphs[0], value, align)
    for cell, caption in zip(table.rows[1].cells, captions):
        _fill(cell.paragraphs[0], caption, size=_CAPTION_SIZE)
    _add_bottom_border(table.rows[0].cells[1])


def _add_title_page(doc, info: TitlePageInfo) -> None:
    for line in _markup_lines(info.university_header):
        _fill(doc.add_paragraph(), line, bold="«" in line or "»" in line, size=10)
    _fill(doc.add_paragraph(), _clean(info.department), space_before=12)

    _fill(doc.add_paragraph(), "ОЦЕНКА РЕФЕРАТА", _LEFT, space_before=48)
    _fill(doc.add_paragraph(), "РУКОВОДИТЕЛЬ", _LEFT)
    _add_signature_table(
        doc,
        (_clean(info.teacher_position), "", _clean(info.teacher_name)),
        ("должность, уч. степень, звание", "подпись, дата", "инициалы, фамилия"),
        _CENTER,
    )

    _fill(doc.add_paragraph(), "РЕФЕРАТ", bold=True, size=16, space_before=60)
    title_lines = _markup_lines(info.lecture_label) or [_clean(info.lecture_title)]
    for index, line in enumerate(title_lines):
        _fill(doc.add_paragraph(), line, bold=True, size=14, space_before=6 if index == 0 else 0)
    _fill(doc.add_paragraph(), f"по дисциплине: {_clean(info.discipline)}", space_before=12)

    _fill(doc.add_paragraph(), "РЕФЕРАТ ВЫПОЛНИЛ", _LEFT, space_before=150)
    _add_signature_table(
        doc,
        (f"СТУДЕНТ гр. № {_clean(info.student_group)}", date.today().strftime("%d.%m.%Y"), _clean(info.student_name)),
        ("", "подпись, дата", "инициалы, фамилия"),
        _LEFT,
    )

    city, comma, year = info.city_year.rpartition(",")
    for index, line in enumerate([city, year] if comma else [info.city_year]):
        _fill(doc.add_paragraph(), _clean(line), space_before=90 if index == 0 else 0)


def _field_char(kind: str):
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    return element


def _add_field(paragraph, instruction: str, placeholder: str) -> None:
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {instruction} "
    for element in (_field_char("begin"), instr_text, _field_char("separate")):
        paragraph.add_run()._r.append(element)
    paragraph.add_run(placeholder)
    paragraph.add_run()._r.append(_field_char("end"))


def _add_table_of_contents(doc) -> None:
    heading = _fill(doc.add_paragraph(), "СОДЕРЖАНИЕ", bold=True)
    _format(heading.paragraph_format, page_break_before=True, space_after=Pt(12))
    _add_field(doc.add_paragraph(), TOC_INSTRUCTION, TOC_PLACEHOLDER)


def _add_sections(doc, sections: list[tuple]) -> None:
    for index, (heading, body, *level) in enumerate(sections):
        style = "Heading 2" if level and level[0] else "Heading 1"
        paragraph = doc.add_paragraph(_clean(heading), style=style)
        if index == 0:
            paragraph.paragraph_format.page_break_before = True
        for chunk in _paragraphs(body):
            doc.add_paragraph(chunk, style=BODY_STYLE)


def _add_sources(doc, sources: list[str]) -> None:
    heading = doc.add_paragraph("СПИСОК ИСТОЧНИКОВ", style="Heading 1")
    _format(heading.paragraph_format, alignment=_CENTER, first_line_indent=Cm(0), page_break_before=True)
    groups = [chunks for chunks in map(_paragraphs, sources) if chunks]
    for number, (first, *rest) in enumerate(groups, start=1):
        doc.add_paragraph(f"{number}. {first}", style=BODY_STYLE)
        for chunk in rest:
            doc.add_paragraph(chunk, style=BODY_STYLE)


def _add_page_number_footer(section) -> None:
    section.different_first_page_header_footer = True
    section.first_page_footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = _CENTER
    _add_field(paragraph, "PAGE", "1")


def _setup_settings(doc) -> None:
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.insert_element_before(update_fields, *_SETTINGS_AFTER_UPDATE_FIELDS)
    # The bundled template is a Word 2010 file; without this Word opens the
    # реферат in compatibility mode.
    for setting in settings.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")


def _set_core_properties(doc, info: TitlePageInfo) -> None:
    properties = doc.core_properties
    properties.author = properties.last_modified_by = _clean(info.student_name)
    properties.title = " ".join(_markup_lines(info.lecture_label)) or _clean(info.lecture_title)
    properties.comments = ""
    properties.revision = 1
    properties.created = properties.modified = datetime.now(timezone.utc)
