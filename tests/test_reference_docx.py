"""Word export of a реферат: the generated .docx is reopened with python-docx
and checked against the ГОСТ 7.32-2017 layout of a real accepted реферат."""
from dataclasses import replace
from datetime import date

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.reference.docx_writer import TOC_PLACEHOLDER, build_reference_docx
from app.reference.pdf_writer import ReferenceContent, TitlePageInfo

HEADER_LINES = [
    "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
    "федеральное государственное автономное образовательное учреждение высшего образования",
    "«ДЕМОНСТРАЦИОННЫЙ УНИВЕРСИТЕТ»",
]

TITLE = TitlePageInfo(
    university_header="<br/>".join(HEADER_LINES),
    department="КАФЕДРА № 00",
    teacher_position="доцент, канд. техн. наук",
    teacher_name="С.С. Сидоров",
    lecture_number=3,
    lecture_title="Коммутация и маршрутизация",
    discipline="Введение в информационные технологии",
    student_group="2395",
    student_name="И.И. Иванов",
    city_year="Санкт-Петербург, 2026",
    lecture_label="к лекциям №&nbsp;3 и 4<br/>«Коммутация &amp; маршрутизация»",
)

SPECIAL_BODY = "Пакет <IP> & кадр Ethernet, причём MTU > 1500 & TTL < 64."

SECTIONS = [
    ("Лекция 3. Коммутация", "", 0),
    (
        "1. Какие методы коммутации бывают?",
        "Коммутация каналов закрепляет физический путь на время сеанса.\n\n"
        f"{SPECIAL_BODY}\n\n\n\nКоммутация пакетов делит данные на части.",
        1,
    ),
    ("2. Что такое СКС & зачем она нужна?", "СКС — структурированная кабельная система.", 1),
    ("3. Итоги", "Коммутация и маршрутизация дополняют друг друга."),
]

SOURCES = [
    "Видеолекция №3 курса «Введение в информационные технологии», 2026.",
    "Сведения об использовании генеративного ИИ.\n\nСистемный промпт: «Напиши реферат».",
]

CENTER = WD_ALIGN_PARAGRAPH.CENTER


def _build(directory, title=TITLE, sections=SECTIONS, sources=SOURCES, font_size=12):
    output = directory / "ЛК3-4_Иванов_2395.docx"
    content = ReferenceContent(title_page=title, sections=sections, sources=sources)
    build_reference_docx(output, content, font_size=font_size)
    return Document(str(output))


@pytest.fixture(scope="module")
def doc(tmp_path_factory):
    return _build(tmp_path_factory.mktemp("docx"))


def _index(doc, text):
    return next(i for i, p in enumerate(doc.paragraphs) if p.text == text)


def _paragraph(doc, text):
    return doc.paragraphs[_index(doc, text)]


def _title_page(doc):
    return doc.paragraphs[: _index(doc, "СОДЕРЖАНИЕ")]


def _effective(paragraph, attribute):
    value = getattr(paragraph.paragraph_format, attribute)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, attribute)
        style = style.base_style
    return value


def _all_text(doc):
    return "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))


def _assert_gost_body(paragraph):
    assert _effective(paragraph, "alignment") == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert round(_effective(paragraph, "first_line_indent").cm, 2) == 1.25
    assert _effective(paragraph, "line_spacing") == 1.5
    assert _effective(paragraph, "space_before") == 0
    assert _effective(paragraph, "space_after") == 0


def test_page_is_a4_with_gost_margins(doc):
    section = doc.sections[0]
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    assert round(section.left_margin.cm, 2) == 3.0
    assert round(section.right_margin.cm, 2) == 1.5
    assert round(section.top_margin.cm, 2) == 2.0
    assert round(section.bottom_margin.cm, 2) == 2.0


def test_times_new_roman_covers_cyrillic(doc):
    normal = doc.styles["Normal"]
    fonts = normal.element.rPr.rFonts
    assert {fonts.get(qn(f"w:{name}")) for name in ("ascii", "hAnsi", "eastAsia", "cs")} == {"Times New Roman"}
    assert fonts.get(qn("w:asciiTheme")) is None
    assert normal.font.size == Pt(12)
    assert normal.font.color.rgb == RGBColor(0, 0, 0)


def test_font_size_parameter_applies_to_text_and_headings(tmp_path):
    doc = _build(tmp_path, font_size=14)
    assert doc.styles["Normal"].font.size == Pt(14)
    assert doc.styles["Heading 1"].font.size == Pt(14)
    assert doc.styles["Heading 2"].font.size == Pt(14)


def test_title_page_contains_form_labels_and_people(doc):
    texts = [p.text for p in _title_page(doc)]
    for expected in (
        "КАФЕДРА № 00",
        "ОЦЕНКА РЕФЕРАТА",
        "РУКОВОДИТЕЛЬ",
        "РЕФЕРАТ",
        "по дисциплине: Введение в информационные технологии",
        "РЕФЕРАТ ВЫПОЛНИЛ",
    ):
        assert expected in texts
    assert _paragraph(doc, "ОЦЕНКА РЕФЕРАТА").alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert _paragraph(doc, "РЕФЕРАТ ВЫПОЛНИЛ").alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_university_header_lines_are_centered_and_name_is_bold(doc):
    header = _title_page(doc)[:3]
    assert [p.text for p in header] == HEADER_LINES
    assert all(p.alignment == CENTER for p in header)
    assert not header[0].runs[0].bold
    assert header[2].runs[0].bold


def test_signature_tables_are_borderless_with_captions(doc):
    teacher, student = doc.tables
    assert [c.text for c in teacher.rows[0].cells] == ["доцент, канд. техн. наук", "", "С.С. Сидоров"]
    assert [c.text for c in teacher.rows[1].cells] == [
        "должность, уч. степень, звание", "подпись, дата", "инициалы, фамилия"
    ]
    assert [c.text for c in student.rows[0].cells] == [
        "СТУДЕНТ гр. № 2395", date.today().strftime("%d.%m.%Y"), "И.И. Иванов"
    ]
    assert [c.text for c in student.rows[1].cells] == ["", "подпись, дата", "инициалы, фамилия"]
    for table in (teacher, student):
        assert table._tbl.tblPr.find(qn("w:tblBorders")) is None
        assert table.style.name == "Normal Table"
        signature_borders = table.rows[0].cells[1]._tc.tcPr.find(qn("w:tcBorders"))
        assert signature_borders.find(qn("w:bottom")).get(qn("w:val")) == "single"
        assert table.rows[0].cells[0]._tc.tcPr.find(qn("w:tcBorders")) is None
        caption = table.rows[1].cells[1].paragraphs[0]
        assert caption.alignment == CENTER
        assert caption.runs[0].font.size == Pt(10)


def test_referat_heading_and_lecture_label(doc):
    referat = _paragraph(doc, "РЕФЕРАТ")
    assert referat.alignment == CENTER and referat.runs[0].bold and referat.runs[0].font.size == Pt(16)
    for line in ("к лекциям № 3 и 4", "«Коммутация & маршрутизация»"):
        paragraph = _paragraph(doc, line)
        assert paragraph.alignment == CENTER
        assert paragraph.runs[0].bold and paragraph.runs[0].font.size == Pt(14)
    assert "Коммутация и маршрутизация" not in [p.text for p in _title_page(doc)]


def test_no_reportlab_markup_leaks_into_document(doc):
    text = _all_text(doc)
    for leftover in ("<br", "&nbsp;", "&amp;", "\xa0"):
        assert leftover not in text


def test_lecture_title_used_when_label_is_empty(tmp_path):
    doc = _build(tmp_path, title=replace(TITLE, lecture_label=""))
    paragraph = _paragraph(doc, "Коммутация и маршрутизация")
    assert paragraph.alignment == CENTER
    assert paragraph.runs[0].bold and paragraph.runs[0].font.size == Pt(14)


def test_city_and_year_are_separate_centered_lines(doc):
    last_two = _title_page(doc)[-2:]
    assert [p.text for p in last_two] == ["Санкт-Петербург", "2026"]
    assert all(p.alignment == CENTER for p in last_two)


def test_city_year_without_comma_stays_one_line(tmp_path):
    doc = _build(tmp_path, title=replace(TITLE, city_year="Санкт-Петербург 2026"))
    texts = [p.text for p in _title_page(doc)]
    assert texts[-1] == "Санкт-Петербург 2026"
    assert "Санкт-Петербург" not in texts


def test_table_of_contents_is_a_word_field_updated_on_open(doc):
    heading = _paragraph(doc, "СОДЕРЖАНИЕ")
    assert heading.alignment == CENTER and heading.runs[0].bold
    assert heading.paragraph_format.page_break_before
    toc = doc.paragraphs[_index(doc, "СОДЕРЖАНИЕ") + 1]
    assert TOC_PLACEHOLDER in toc.text
    instructions = [el.text.strip() for el in toc._p.iter(qn("w:instrText"))]
    assert instructions == ['TOC \\o "1-2" \\h \\z \\u']
    assert [el.get(qn("w:fldCharType")) for el in toc._p.iter(qn("w:fldChar"))] == ["begin", "separate", "end"]

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    assert update_fields is not None and update_fields.get(qn("w:val")) == "true"
    tags = [child.tag for child in settings]
    assert tags.index(qn("w:updateFields")) < tags.index(qn("w:compat"))
    modes = [s.get(qn("w:val")) for s in settings.iter(qn("w:compatSetting")) if s.get(qn("w:name")) == "compatibilityMode"]
    assert modes == ["15"]

    for level in (1, 2):
        toc_style = doc.styles[f"toc {level}"]
        assert toc_style.builtin, "Word ignores a custom 'toc N' style when it builds the TOC"
        assert toc_style.paragraph_format.line_spacing == 1.5


def test_section_headings_use_builtin_heading_styles_per_level(doc):
    headings = [(p.text, p.style.name) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [
        ("Лекция 3. Коммутация", "Heading 1"),
        ("1. Какие методы коммутации бывают?", "Heading 2"),
        ("2. Что такое СКС & зачем она нужна?", "Heading 2"),
        ("3. Итоги", "Heading 1"),
        ("СПИСОК ИСТОЧНИКОВ", "Heading 1"),
    ]


def test_group_heading_with_empty_body_starts_the_text_on_a_new_page(doc):
    index = _index(doc, "Лекция 3. Коммутация")
    assert doc.paragraphs[index].paragraph_format.page_break_before
    assert doc.paragraphs[index + 1].text == "1. Какие методы коммутации бывают?"


def test_body_paragraphs_split_on_blank_lines_and_skip_empty_ones(doc):
    start = _index(doc, "1. Какие методы коммутации бывают?")
    assert [p.text for p in doc.paragraphs[start + 1 : start + 4]] == [
        "Коммутация каналов закрепляет физический путь на время сеанса.",
        SPECIAL_BODY,
        "Коммутация пакетов делит данные на части.",
    ]
    for paragraph in doc.paragraphs[start + 1 : start + 4]:
        assert paragraph.style.name == "Body Text"
        _assert_gost_body(paragraph)


def test_heading_styles_are_restyled_to_gost(doc):
    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        font = style.font
        assert font.name == "Times New Roman"
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == "Times New Roman"
        assert style.element.rPr.rFonts.get(qn("w:asciiTheme")) is None
        assert font.size == Pt(12) and font.bold and font.italic is False
        assert font.color.rgb == RGBColor(0, 0, 0) and font.color.theme_color is None
        paragraph_format = style.paragraph_format
        assert paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert round(paragraph_format.first_line_indent.cm, 2) == 1.25
        assert paragraph_format.line_spacing == 1.5
        assert paragraph_format.space_before == Pt(12) and paragraph_format.space_after == Pt(6)
        assert paragraph_format.keep_with_next
    assert doc.styles["Heading 2"].paragraph_format.left_indent == 0


def test_plain_text_markup_characters_survive_round_trip(doc):
    texts = [p.text for p in doc.paragraphs]
    assert SPECIAL_BODY in texts
    assert "2. Что такое СКС & зачем она нужна?" in texts


def test_sources_heading_is_centered_heading_1_with_numbered_sources(doc):
    index = _index(doc, "СПИСОК ИСТОЧНИКОВ")
    heading = doc.paragraphs[index]
    assert heading.style.name == "Heading 1"
    assert heading.alignment == CENTER
    assert heading.paragraph_format.first_line_indent == 0
    assert heading.paragraph_format.page_break_before

    sources = doc.paragraphs[index + 1 :]
    assert [p.text for p in sources] == [
        "1. Видеолекция №3 курса «Введение в информационные технологии», 2026.",
        "2. Сведения об использовании генеративного ИИ.",
        "Системный промпт: «Напиши реферат».",
    ]
    for paragraph in sources:
        _assert_gost_body(paragraph)


def test_empty_sources_are_not_numbered(tmp_path):
    doc = _build(tmp_path, sources=["", "Первый источник.", "  \n\n  ", "Второй источник."])
    index = _index(doc, "СПИСОК ИСТОЧНИКОВ")
    assert [p.text for p in doc.paragraphs[index + 1 :]] == ["1. Первый источник.", "2. Второй источник."]


def test_footer_page_number_hidden_on_title_page(doc):
    section = doc.sections[0]
    assert section.different_first_page_header_footer
    footer = section.footer
    assert [el.text.strip() for el in footer._element.iter(qn("w:instrText"))] == ["PAGE"]
    assert footer.paragraphs[0].alignment == CENTER

    first_page_footer = section.first_page_footer
    assert not first_page_footer.is_linked_to_previous
    assert list(first_page_footer._element.iter(qn("w:instrText"))) == []
    assert all(not p.text for p in first_page_footer.paragraphs)


def test_core_properties_name_the_student_not_the_library(doc):
    properties = doc.core_properties
    assert properties.author == "И.И. Иванов"
    assert properties.last_modified_by == "И.И. Иванов"
    assert "python-docx" not in (properties.comments or "")
    assert properties.title == "к лекциям № 3 и 4 «Коммутация & маршрутизация»"
